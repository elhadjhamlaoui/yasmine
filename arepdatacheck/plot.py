import os
import unicodedata
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

folder = r"C:\Users\bompoloe\Desktop\Projets\CFL LUX\AUDITS\2026_09_17\PLB"

# Fichier genere par export_parametres_applicables.py (colonnes Categorie;Parametre)
applicability_file = os.path.join(folder, "parametres_applicables.csv")

CATEGORY_COLUMNS = ["categorie", "category", "famille"]
PARAMETER_COLUMNS = ["parametre", "parameter", "nom_parametre"]


def norm(value):
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.strip().casefold()


def find_column(df, candidates):
    for col in df.columns:
        if norm(col) in candidates:
            return col
    return None


def load_applicability(path):
    if not os.path.isfile(path):
        print("ATTENTION : {} introuvable, aucun filtrage applique.".format(path))
        return None
    ref = pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig")
    applicable = {}
    for cat, param in zip(ref.iloc[:, 0], ref.iloc[:, 1]):
        applicable.setdefault(norm(cat), set()).add(norm(param))
    return applicable


def filter_not_applicable(df, applicable, name):
    """Retire les parametres non coches pour la categorie.
    - format long (colonnes Categorie + Parametre) : la ligne est supprimee
    - format large (Categorie en lignes, parametres en colonnes) : la cellule est videe
    """
    if not applicable:
        return df
    cat_col = find_column(df, CATEGORY_COLUMNS)
    if cat_col is None:
        return df
    param_col = find_column(df, PARAMETER_COLUMNS)

    if param_col is not None:
        keep = [
            norm(param) in applicable.get(norm(cat), set())
            for cat, param in zip(df[cat_col], df[param_col])
        ]
        removed = len(keep) - sum(keep)
        if removed:
            print("{} : {} ligne(s) non applicable(s) ignoree(s)".format(name, removed))
        return df[keep].reset_index(drop=True)

    df = df.copy()
    blanked = 0
    for i, cat in enumerate(df[cat_col]):
        allowed = applicable.get(norm(cat))
        if allowed is None:
            continue
        for col in df.columns:
            if col != cat_col and norm(col) not in allowed:
                if not pd.isna(df.at[i, col]):
                    blanked += 1
                df.at[i, col] = None
    if blanked:
        print("{} : {} cellule(s) non applicable(s) videe(s)".format(name, blanked))
    return df


applicability = load_applicability(applicability_file)

files = [
    f for f in os.listdir(folder)
    if os.path.isfile(os.path.join(folder, f)) and f.lower().endswith(".csv")
    and f != "parametres_applicables.csv"
]

output_excel = os.path.join(
    r"C:\Users\bompoloe\Desktop\Projets\CFL LUX\AUDITS",
    "rapport_analyse_données_MN_PLB.xlsx"
)

# Cree un fichier vide au depart
pd.DataFrame().to_excel(output_excel, sheet_name="Parametres", index=False)

wb = load_workbook(output_excel)
ws = wb["Parametres"]

vertical_gap = 3
right_side_sources = ["Parametre_manquant_data", "score_pondré_parametres_data"]
left_column = 1
right_column = 16
left_start_row = 1
right_start_row = 4
next_row_by_column = {
    left_column: left_start_row,
    right_column: right_start_row,
}

for filename in files:
    path = os.path.join(folder, filename)
    print("reading")
    df = pd.read_csv(
        path,
        sep=None,
        engine="python",
        encoding="utf-8",
    )
    df = filter_not_applicable(df, applicability, filename)

    table_name = os.path.splitext(filename)[0]
    start_column = right_column if  right_side_sources[0] in table_name or right_side_sources[1] in table_name  else left_column
    base_row = next_row_by_column[start_column]
    start_row = base_row if start_column == right_column and base_row == right_start_row else base_row + vertical_gap

    title = filename
    if "_détaché_" in filename:
        title = filename.split("_détaché_")[1].replace(".csv", "")
    elif "_dÃ©tachÃ©_" in filename:
        title = filename.split("_dÃ©tachÃ©_")[1].replace(".csv", "")
    else:
        title = os.path.splitext(filename)[0]

    title_cell = ws.cell(row=start_row, column=start_column, value=title)
    title_cell.fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
    ws.merge_cells(
        start_row=start_row,
        start_column=start_column,
        end_row=start_row,
        end_column=start_column + len(df.columns) - 1
    )

    for c, colname in enumerate(df.columns):
        header = ws.cell(row=start_row + 1, column=start_column + c, value=colname)
        header.fill = PatternFill(start_color="c7c7c7", end_color="c7c7c7", fill_type="solid")

    for r in range(len(df)):
        for c in range(len(df.columns)):
            value = df.iloc[r, c]
            cell = ws.cell(row=start_row + 2 + r, column=start_column + c, value=None if pd.isna(value) else value)
            if cell.value == 0:
                cell.fill = PatternFill(start_color="f58c9b", end_color="f58c9b", fill_type="solid")

    next_row_by_column[start_column] = start_row + len(df) + 1

wb.save(output_excel)
print("Fichier généré :", output_excel)
